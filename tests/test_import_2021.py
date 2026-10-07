import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from pinhook.obs_legacy_excel import normalize
from pinhook.import_2021 import import_files


class Import2021Test(unittest.TestCase):
    def test_2022_spring_xlsx_statuses_and_birth_date(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'obs_spring22_results.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.append(['Hip', 'Name', 'Color', 'Sex', 'Foal Date', 'Sire', 'Dam', 'Damsire',
                          'Extra', 'Consignor', 'A', 'B', 'C', 'D', 'E', 'Buyer', 'Price'])
            sheet.append([1, '', 'B', 'C', datetime(2020, 3, 5), 'Sire', 'Dam', 'Damsire',
                          '', 'Agent', '', '', '', '', '', 'Buyer', 75000])
            sheet.append([2, '', 'B', 'F', datetime(2020, 4, 7), 'Sire', 'Dam2', 'Damsire',
                          '', 'Agent', '', '', '', '', '', 42000, 'Not Sold'])
            book.save(path)
            rows, excluded = normalize(path, 'spring22')
            self.assertEqual(len(rows), 2)
            self.assertEqual(rows[0]['foaling_date'], '2020-03-05')
            self.assertEqual(rows[0]['status'], 'sold')
            self.assertEqual(rows[1]['reported_bid_usd'], '42000.00')
            self.assertEqual(rows[1]['price_usd'], '')
            self.assertEqual(excluded, [])

    def test_archive_header_labels_as_published_2020_to_2022(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'obs_spring22_results.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.append(['OBS SPRING SALE OF TWO-YEAR-OLDS IN TRAINING'])
            sheet.append(['hip#', 'Name', 'Color', 'Sex', 'Foal Date', 'Sire', 'Dam', 'Dam Sire',
                          'Sort By Dam', 'Consignor', 'Area ID', 'Barn', 'work time', 'set', 'day', 'Buyer ', 'Price ', 'PS'])
            sheet.append([1, '', 'B', 'C', datetime(2020, 3, 5), 'Sire', 'Dam', 'Damsire',
                          '', 'Agent', '', '', '', '', '', 'Buyer', 75000])
            book.save(path)
            rows, _ = normalize(path, 'spring22')
            self.assertEqual([r['status'] for r in rows], ['sold'])

    def test_2020_placeholder_and_racing_age_rows_are_excluded(self):
        header = ['hip#', 'Name', 'Color', 'Sex', 'Foal Date', 'Sire', 'Dam', 'Dam Sire',
                  'Sort By Dam', 'Consignor', 'Area ID', 'Barn', 'work time', 'set', 'day', 'Buyer ', 'Price ']
        cases = [('march20', [87, '.', '.', '.', '.', '.', '.', '.', '.', '.', '.', '.', 'out', '', '', 'Withdrawn', 'Out']),
                 ('july20', [991, 'Law', 'B', 'F', datetime(2017, 1, 13), 'Constitution', 'Dam', 'Damsire',
                             '', 'Agent', '', '', '', '', '', 'Withdrawn', 'Out'])]
        for sale, row in cases:
            with tempfile.TemporaryDirectory() as directory, self.subTest(sale=sale):
                path = Path(directory) / f'obs_{sale}_results.xlsx'
                book = Workbook()
                book.active.append(header)
                book.active.append(row)
                book.save(path)
                self.assertEqual(normalize(path, sale), ([], [row[0]]))

    def test_incomplete_downloads_do_not_build_cohort(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self.assertFalse(import_files(root / 'raw', root / 'work'))
            self.assertFalse((root / 'work' / 'cohort_2021_to_2022').exists())


if __name__ == '__main__':
    unittest.main()
